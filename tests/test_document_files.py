import io

import openpyxl
from docx import Document as DocxDocument
from pypdf import PdfReader

from core import document_files


def test_build_docx_contains_title_and_paragraphs():
    data = document_files.build_docx("Письмо клиенту", "Первый абзац.\n\nВторой абзац.")

    doc = DocxDocument(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Письмо клиенту" in text
    assert "Первый абзац." in text
    assert "Второй абзац." in text


def test_build_xlsx_creates_rows_from_pipe_delimited_text():
    data = document_files.build_xlsx("Статья|Сумма\nАренда|30000\nЗарплата|80000")

    workbook = openpyxl.load_workbook(io.BytesIO(data))
    sheet = workbook.active
    rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    assert rows == [("Статья", "Сумма"), ("Аренда", "30000"), ("Зарплата", "80000")]


def test_build_pdf_produces_readable_cyrillic_text():
    data = document_files.build_pdf("Счёт на оплату", "Уважаемый клиент,\n\nОплатите счёт до пятницы.")

    reader = PdfReader(io.BytesIO(data))
    text = "".join(page.extract_text() for page in reader.pages)
    assert "Счёт на оплату" in text
    assert "Оплатите счёт до пятницы" in text


def test_build_document_file_routes_by_format():
    docx_bytes, docx_name = document_files.build_document_file("docx", "Письмо/тест", "текст")
    xlsx_bytes, xlsx_name = document_files.build_document_file("xlsx", "Таблица", "a|b")
    pdf_bytes, pdf_name = document_files.build_document_file("pdf", "Документ", "текст")

    assert docx_name == "Письмотест.docx"
    assert xlsx_name == "Таблица.xlsx"
    assert pdf_name == "Документ.pdf"
    assert docx_bytes and xlsx_bytes and pdf_bytes


def test_build_document_file_raises_on_unknown_format():
    import pytest

    with pytest.raises(ValueError):
        document_files.build_document_file("txt", "заголовок", "содержимое")
